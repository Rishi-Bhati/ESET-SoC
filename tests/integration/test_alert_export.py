"""
The Alerts view's "Convert to Spreadsheet" export (GET /dashboard/api/alerts/export,
src/services/alert_export.py): the complete alert history, one row per alert,
every detail-view field, regardless of the table's window, filters or paging.
"""
import asyncio
import io
import json
import os
import re
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.config import settings
from src.services.alert_export import EXCEL_CELL_MAX, _SheetWriter, _flatten
from src.storage import job_store
from src.utils.jst_time import JST
from ai_fakes import FakeProvider, default_responder

AUTH = {"Authorization": "Bearer test_token"}
EXPORT = "/dashboard/api/alerts/export"
SEVERITIES = ("LOW", "MEDIUM", "HIGH", "CRITICAL")


def _payload(i: int, **extra) -> dict:
    base = {
        "alert_id": f"export-{i:03d}",
        "occurred_at": f"2026-09-20T{i % 24:02d}:05:00Z",
        "severity": SEVERITIES[i % 4],
        "detection_name": f"Win32/Export.Test.{i}",
        "endpoint_name": f"HOST-{i:03d}",
        "threat_handled": i % 2 == 0,
        "isolation_status": False,
        "file_hash": f"{i:064x}",
        "raw_content": f"Detection number {i}",
    }
    base.update(extra)
    return base


def _ingest(client: TestClient, i: int, route: str = "/webhook/eset", **extra) -> str:
    res = client.post(route, headers=AUTH, json=_payload(i, **extra))
    assert res.status_code == 200, res.text
    assert res.json()["status"] == "queued"
    return res.json()["correlation_id"]


def _export(client: TestClient, **kwargs):
    res = client.get(EXPORT, **kwargs)
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert re.fullmatch(r'attachment; filename="soc-lite-alerts-\d{8}T\d{6}Z\.xlsx"',
                        res.headers["content-disposition"])
    return res, load_workbook(io.BytesIO(res.content))


def _rows(wb, sheet="Alerts") -> list[dict]:
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, values)) for values in ws.iter_rows(min_row=2, values_only=True)]


@pytest.fixture
def recipients(monkeypatch):
    monkeypatch.setattr(settings, "client_notification_emails", "client@example.com")
    monkeypatch.setattr(settings, "cthree_notification_emails", "cthree@example.com")
    monkeypatch.setattr(settings, "internal_notification_emails", "internal@example.com")
    monkeypatch.setattr(settings, "engineer_notification_emails", "eng@example.com")


@pytest.fixture
def sixty_alerts(client: TestClient, recipients) -> list[str]:
    """60 alerts through the real ingest path: webhook and syslog, every risk level."""
    return [_ingest(client, i, "/webhook/syslog" if i % 5 == 0 else "/webhook/eset") for i in range(60)]


# ------------------------------------------------------------ completeness

def test_all_sixty_alerts_are_exported_once_each(client: TestClient, sixty_alerts):
    res, wb = _export(client)
    rows = _rows(wb)
    assert res.headers["x-alert-count"] == "60"
    assert len(rows) == 60
    assert sorted(r["correlation_id"] for r in rows) == sorted(sixty_alerts), "every alert, no duplicates"
    assert {r["record"] for r in rows} == {"job + result"}
    assert {r["risk_level"] for r in rows} <= set(SEVERITIES) and None not in {r["risk_level"] for r in rows}
    assert {r["job.source"] for r in rows} == {"WEBHOOK", "SYSLOG"}
    info = dict(wb["Export info"].iter_rows(values_only=True))
    assert info["Alerts"] == 60


def test_new_alerts_are_in_the_next_export_without_duplicates(client: TestClient, sixty_alerts):
    _, first = _export(client)
    assert len(_rows(first)) == 60

    new_ids = [_ingest(client, 100), _ingest(client, 101, "/webhook/syslog")]
    # A resend of an existing alert is a duplicate at ingest, never a second row.
    assert client.post("/webhook/eset", headers=AUTH, json=_payload(3)).json()["status"] == "duplicate"

    res, second = _export(client)
    ids = [r["correlation_id"] for r in _rows(second)]
    assert res.headers["x-alert-count"] == "62"
    assert len(ids) == len(set(ids)) == 62
    assert set(ids) == set(sixty_alerts) | set(new_ids)


def test_export_ignores_the_dashboard_window_filters_and_paging(client: TestClient):
    # The Alerts table loads the newest 500 jobs; the export must not stop there.
    async def create_jobs():
        for i in range(600):
            await job_store.create_job(f"bulk-{i:04d}", "WEBHOOK", {"alert_id": f"bulk-{i}"})
    asyncio.run(create_jobs())
    assert len(client.get("/dashboard/api/jobs?limit=500").json()["jobs"]) == 500

    res, wb = _export(client, params={"status": "SUCCESS", "limit": 10, "offset": 50, "page": 3})
    rows = _rows(wb)
    assert len(rows) == 600 == int(res.headers["x-alert-count"])
    assert {r["record"] for r in rows} == {"job only (no result yet)"}, "unfinished alerts are included too"
    assert {r["job.status"] for r in rows} == {"PENDING"}


def test_a_retried_alert_is_still_one_row(client: TestClient, recipients, monkeypatch):
    monkeypatch.setattr(FakeProvider, "responder", staticmethod(lambda request: TimeoutError("down")))
    cid = _ingest(client, 7)
    assert client.get(f"/dashboard/api/jobs/{cid}").json()["job"]["status"] == "PARTIAL"

    monkeypatch.setattr(FakeProvider, "responder", staticmethod(default_responder))
    assert client.post(f"/dashboard/api/jobs/{cid}/retry").status_code == 200
    assert client.get(f"/dashboard/api/jobs/{cid}").json()["job"]["status"] == "SUCCESS"
    # index.json now lists the alert twice (one record per pipeline run) ...
    with open(os.path.join(settings.output_dir, "index.json"), encoding="utf-8") as f:
        assert [r["correlation_id"] for r in json.load(f)].count(cid) == 2

    # ... but the export is keyed by correlation_id and reflects the latest run.
    rows = _rows(_export(client)[1])
    assert [r["correlation_id"] for r in rows] == [cid]
    assert rows[0]["job.status"] == "SUCCESS" and rows[0]["ai_run.status"] == "SUCCESS"


def test_a_result_file_without_a_job_row_is_not_lost(client: TestClient, sixty_alerts):
    orphan_source = os.path.join(settings.output_dir, f"{sixty_alerts[0]}.json")
    with open(orphan_source, encoding="utf-8") as f:
        orphan = json.load(f)
    orphan["correlation_id"] = "orphan-0001"
    with open(os.path.join(settings.output_dir, "orphan-0001.json"), "w", encoding="utf-8") as f:
        json.dump(orphan, f)

    rows = {r["correlation_id"]: r for r in _rows(_export(client)[1])}
    assert len(rows) == 61
    assert rows["orphan-0001"]["record"] == "result file only (no job row)"
    assert rows["orphan-0001"]["normalized_alert.detection_name"] == orphan["normalized_alert"]["detection_name"]


def test_an_empty_history_exports_a_valid_workbook(client: TestClient):
    res, wb = _export(client)
    assert res.headers["x-alert-count"] == "0"
    assert wb.sheetnames == ["Alerts", "Risk factors", "Emails", "Export info"]
    assert "correlation_id" in [c.value for c in wb["Alerts"][1]]


# ---------------------------------------------------------------- details

def test_every_detail_view_field_is_a_column_with_its_value(client: TestClient, sixty_alerts):
    cid = sixty_alerts[2]
    detail = client.get(f"/dashboard/api/jobs/{cid}").json()
    job, result = detail["job"], detail["result"]
    row = next(r for r in _rows(_export(client)[1]) if r["correlation_id"] == cid)

    expected: dict = {}
    special = {"correlation_id", "received_at", "processed_at", "risk_factors", "notifications"}
    _flatten({k: v for k, v in result.items() if k not in special}, "", expected)
    for column in [c for c in expected if c.startswith("normalized_alert.raw_payload")]:
        expected.pop(column)                                  # exported as raw_request (JSON)
    expected.pop("normalized_alert.occurred_at")              # exported as-received + JST
    for column, value in expected.items():
        assert column in row, f"missing column {column}"
        assert row[column] == (value if value != "" else None), column

    assert row["normalized_alert.occurred_at (as received)"] == "2026-09-20T02:05:00Z"
    assert row["normalized_alert.occurred_at (JST)"] == datetime(2026, 9, 20, 11, 5)
    created_jst = datetime.fromtimestamp(job["created_at"], tz=timezone.utc).astimezone(JST).replace(tzinfo=None)
    # Excel stores date-times to the millisecond.
    assert abs(row["job.created_at (JST)"] - created_jst) <= timedelta(milliseconds=1)
    assert row["job.status"] == job["status"] and row["job.source"] == job["source"]
    assert json.loads(row["raw_request (JSON)"]) == job["raw_payload"]["raw_payload"]
    for f in result["risk_factors"]:
        assert f["rule"] in row["risk_factors"] and f["detail"] in row["risk_factors"]
    # The AI's notification texts, as on the detail view's tabs.
    for field in ("alert_summary_ja", "email_body_ja", "engineer_summary_en", "backlog_comment_ja"):
        assert row[f"ai_output.{field}"] == result["ai_output"][field]
    assert row["ai_run.request_id"] == "req_mock_0001"
    assert row["threat_intel.virustotal.status"] == result["threat_intel"]["virustotal"]["status"]


def test_risk_factor_and_email_sheets_cover_every_alert(client: TestClient, sixty_alerts):
    _, wb = _export(client)
    factors, emails = _rows(wb, "Risk factors"), _rows(wb, "Emails")
    assert {r["correlation_id"] for r in factors} == set(sixty_alerts)
    # Four audiences per successful alert, each listed once.
    assert len(emails) == 240
    assert len({r["email_id"] for r in emails}) == 240
    assert {r["notification_type"] for r in emails} == {"CLIENT_JA", "CTHREE_JA", "INTERNAL_JA", "ENGINEER_EN"}
    assert {r["recipients"] for r in emails if r["notification_type"] == "CLIENT_JA"} == {"client@example.com"}


# --------------------------------------------------------------- security

def test_alert_content_is_stored_as_text_never_as_a_formula(client: TestClient):
    formula = '=HYPERLINK("http://evil.example","click")'
    cid = _ingest(client, 1, detection_name=formula, endpoint_name="+cmd|' /C calc'!A0")
    res = client.get(EXPORT)
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["Alerts"]
    header = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == cid)
    cell = row[header.index("normalized_alert.detection_name")]
    assert cell.value == formula and cell.data_type == "s"
    assert b"<f>" not in _sheet_xml(res.content), "no cell may be written as a formula"


def _sheet_xml(content: bytes) -> bytes:
    import zipfile
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        return b"".join(z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/"))


def test_export_requires_the_dashboard_key(client: TestClient, monkeypatch):
    monkeypatch.setattr(settings, "dashboard_access_key", "export-key")
    assert client.get(EXPORT).status_code == 401
    assert client.get(EXPORT, headers={"X-Dashboard-Key": "wrong"}).status_code == 401
    assert client.get(EXPORT, headers={"X-Dashboard-Key": "export-key"}).status_code == 200


def test_oversized_and_illegal_values_are_made_storable():
    writer = _SheetWriter()
    long_text = writer.cell_value("x" * (EXCEL_CELL_MAX + 500))
    assert len(long_text) == EXCEL_CELL_MAX and "truncated" in long_text and writer.truncated_cells == 1
    assert writer.cell_value("a\x00b\x1bc") == "a�b�c"


# ------------------------------------------------------------------- UI

STATIC = os.path.join(os.path.dirname(__file__), "..", "..", "static")


def test_the_alerts_view_has_one_convert_button():
    with open(os.path.join(STATIC, "dashboard.html"), encoding="utf-8") as f:
        html = f.read()
    alerts_view = html.split('id="view-alerts"', 1)[1].split("</section>", 1)[0]
    assert alerts_view.count('id="alertsExport"') == 1
    assert html.count('id="alertsExport"') == 1
    assert 'data-i18n="btn_export_alerts_xlsx"' in alerts_view
    with open(os.path.join(STATIC, "dashboard.js"), encoding="utf-8") as f:
        assert '"/alerts/export"' in f.read()


def test_export_strings_are_translated_in_every_language():
    with open(os.path.join(STATIC, "i18n.js"), encoding="utf-8") as f:
        source = f.read()
    languages = re.findall(r"^  (\w+): \{$", source, flags=re.MULTILINE)
    assert set(languages) >= {"ja", "en"}
    blocks = re.split(r"^  \w+: \{$", source, flags=re.MULTILINE)[1:]
    for lang, block in zip(languages, blocks):
        for key in ("btn_export_alerts_xlsx", "btn_exporting_alerts",
                    "toast_alerts_exported", "toast_alerts_export_failed"):
            assert re.search(rf"\b{key}:", block), f"{key} missing for {lang}"
