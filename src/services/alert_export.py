"""
The Alerts view's "Convert to Spreadsheet" export: every alert the platform
has recorded, as one .xlsx workbook.

Source of truth is the same data the dashboard reads:

  * the `jobs` table — one row per accepted alert, written at ingest for the
    webhook and syslog paths alike (src/api/webhook.py), so an alert is in the
    next export the moment it is received, finished or not;
  * each alert's result file (output/alerts/{correlation_id}.json) — what the
    alert detail view shows: normalized alert, risk decision, threat intel,
    AI output and AI run record;
  * the `email_deliveries` table — the Timeline's EMAIL stage.

Everything is read unpaged, and rows are keyed by correlation_id, so the
export is never limited by the dashboard's table window, filters or page size,
and an alert retried (which re-appends to index.json) still appears once.
A result file whose job row is gone is exported too, marked as such, rather
than silently dropped.

Every field of the result record is flattened into its own column, so a field
added to the model later appears without changing this module. Timestamps are
written as Excel date-times in JST (src/utils/jst_time.py); the source's own
occurred_at string is kept alongside, since it is not always convertible.

Alert content is attacker-influenced, so every string is stored as text —
never as a formula — and capped at Excel's per-cell limit.
"""
import asyncio
import io
import json
import os
import re
from datetime import datetime, timezone
from typing import Any

import structlog
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.config import settings
from src.storage import delivery_store, job_store
from src.utils.jst_time import CONVERTED, JST, parse_alert_time

logger = structlog.get_logger(__name__)

EXCEL_CELL_MAX = 32_767
_DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
_CORRELATION_ID_RE = re.compile(r"[A-Za-z0-9-]{1,64}")

RECORD_COMPLETE = "job + result"
RECORD_IN_PROGRESS = "job only (no result yet)"
RECORD_RESULT_ONLY = "result file only (no job row)"

# Lead columns of the Alerts sheet, in this order; every other field follows in
# the order first seen, and the raw request comes last (it is the widest).
_LEADING_COLUMNS = (
    "correlation_id", "record", "job.status", "pipeline_status", "risk_level",
    "normalized_alert.severity", "normalized_alert.detection_name", "normalized_alert.endpoint_name",
    "job.source", "normalized_alert.occurred_at (JST)", "normalized_alert.occurred_at (as received)",
    "received_at (JST)", "processed_at (JST)", "job.created_at (JST)", "job.updated_at (JST)",
    "job.error", "error", "risk_rationale", "risk_factors", "notifications", "email_deliveries",
)
_TRAILING_COLUMNS = ("raw_request (JSON)",)

# Result fields given their own treatment below instead of plain flattening.
# normalized_alert.raw_payload is the original request, exported once as
# raw_request (JSON) — the same thing the dashboard's Raw button shows.
_SPECIAL_RESULT_FIELDS = {
    "correlation_id", "received_at", "processed_at", "risk_factors", "notifications",
}


# ------------------------------------------------------------------ values

def _epoch_jst(value: Any) -> datetime | None:
    if not isinstance(value, (int, float)):
        return None
    return datetime.fromtimestamp(value, tz=timezone.utc).astimezone(JST).replace(tzinfo=None)


def _iso_jst(value: Any) -> datetime | str | None:
    """A timestamp string as a JST date-time, or the string as given if its zone is unknown."""
    if value in (None, ""):
        return None
    parsed = parse_alert_time(value)
    return parsed.jst.replace(tzinfo=None) if parsed.status == CONVERTED else parsed.en


def _is_scalar(value: Any) -> bool:
    return value is None or isinstance(value, (str, int, float, bool))


def _flatten(value: Any, prefix: str, out: dict[str, Any]) -> None:
    """Nested dicts become dotted column names; lists become one cell each."""
    if isinstance(value, dict):
        if not value and prefix:
            out[prefix] = None
        for key, item in value.items():
            _flatten(item, f"{prefix}.{key}" if prefix else str(key), out)
    elif isinstance(value, list):
        if all(_is_scalar(item) for item in value):
            out[prefix] = "\n".join("" if item is None else str(item) for item in value)
        else:
            out[prefix] = json.dumps(value, ensure_ascii=False, default=str)
    else:
        out[prefix] = value


def _original_request(job: dict[str, Any] | None, result: dict[str, Any] | None) -> Any:
    """The original submitted JSON, as the dashboard's Raw view shows it
    (static/dashboard.js: originalRequestPayload)."""
    if job:
        outer = job.get("raw_payload") or {}
        inner = outer.get("raw_payload") if isinstance(outer, dict) else None
        return inner if isinstance(inner, dict) and inner else outer
    return ((result or {}).get("normalized_alert") or {}).get("raw_payload") or {}


# ------------------------------------------------------------------- rows

def _alert_row(job: dict[str, Any] | None, result: dict[str, Any] | None,
               deliveries: list[dict[str, Any]]) -> dict[str, Any]:
    cid = (job or result or {}).get("correlation_id")
    row: dict[str, Any] = {
        "correlation_id": cid,
        "record": RECORD_COMPLETE if job and result else RECORD_IN_PROGRESS if job else RECORD_RESULT_ONLY,
    }
    if job:
        row["job.status"] = job.get("status")
        row["job.source"] = job.get("source")
        row["job.created_at (JST)"] = _epoch_jst(job.get("created_at"))
        row["job.updated_at (JST)"] = _epoch_jst(job.get("updated_at"))
        row["job.error"] = job.get("error")

    if result:
        row["received_at (JST)"] = _iso_jst(result.get("received_at"))
        row["processed_at (JST)"] = _iso_jst(result.get("processed_at"))
        rest = {k: v for k, v in result.items() if k not in _SPECIAL_RESULT_FIELDS}
        normalized = dict(rest.get("normalized_alert") or {})
        normalized.pop("raw_payload", None)
        occurred = normalized.pop("occurred_at", None)
        rest["normalized_alert"] = normalized
        _flatten(rest, "", row)
        row["normalized_alert.occurred_at (as received)"] = occurred
        row["normalized_alert.occurred_at (JST)"] = _iso_jst(occurred) if occurred != "UNKNOWN" else None
        row["risk_factors"] = "\n".join(
            f"{f.get('rule')} [{f.get('effect')}]: {f.get('detail')}" for f in result.get("risk_factors") or [])
        row["notifications"] = "\n".join(
            f"{n.get('notification_type')} ({n.get('kind')}): {', '.join(n.get('recipients') or [])}"
            for n in result.get("notifications") or [])

    row["email_deliveries"] = "\n".join(
        f"{d['notification_type']}: {d['status']} (attempts {d['attempts']})" for d in deliveries)
    row["raw_request (JSON)"] = json.dumps(_original_request(job, result), ensure_ascii=False, default=str)
    return row


def _email_rows(cid: str, result: dict[str, Any] | None,
                deliveries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per staged notification, joined to its handoff record by email_id."""
    by_id = {d["email_id"]: d for d in deliveries}
    rows = []
    for n in (result or {}).get("notifications") or []:
        rows.append(_email_row(cid, n, by_id.pop(n.get("email_id"), None)))
    for d in by_id.values():          # a handoff record whose notification is not in the result file
        rows.append(_email_row(cid, None, d))
    return rows


def _email_row(cid: str, notification: dict[str, Any] | None, delivery: dict[str, Any] | None) -> dict[str, Any]:
    n, d = notification or {}, delivery or {}
    return {
        "correlation_id": cid,
        "email_id": n.get("email_id") or d.get("email_id"),
        "notification_type": n.get("notification_type") or d.get("notification_type"),
        "kind": n.get("kind"),
        "recipients": "\n".join(n.get("recipients") or d.get("recipients") or []),
        "subject": d.get("subject"),
        "delivery_status": d.get("status"),
        "attempts": d.get("attempts"),
        "remote_id": d.get("remote_id"),
        "error": d.get("error"),
        "created_at (JST)": _epoch_jst(d.get("created_at")),
        "updated_at (JST)": _epoch_jst(d.get("updated_at")),
    }


def _risk_factor_rows(cid: str, result: dict[str, Any] | None) -> list[dict[str, Any]]:
    return [
        {"correlation_id": cid, "order": i, "rule": f.get("rule"), "effect": f.get("effect"), "detail": f.get("detail")}
        for i, f in enumerate((result or {}).get("risk_factors") or [], start=1)
    ]


# ------------------------------------------------------------------ reading

def _read_result(correlation_id: str) -> dict[str, Any] | None:
    if not _CORRELATION_ID_RE.fullmatch(correlation_id or ""):
        return None
    path = os.path.join(settings.output_dir, f"{correlation_id}.json")
    if not os.path.exists(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        logger.warning("alert_export_result_unreadable", correlation_id=correlation_id, error=str(e))
        return None


def _result_ids_on_disk() -> list[str]:
    try:
        names = os.listdir(settings.output_dir)
    except FileNotFoundError:
        return []
    return sorted(
        name[:-5] for name in names
        if name.endswith(".json") and name != "index.json" and _CORRELATION_ID_RE.fullmatch(name[:-5])
    )


def _collect(jobs: list[dict[str, Any]], deliveries: list[dict[str, Any]]) -> list[tuple]:
    """(job, result, deliveries) per alert: every job, then any result file without one."""
    deliveries_by_cid: dict[str, list[dict[str, Any]]] = {}
    for d in deliveries:
        deliveries_by_cid.setdefault(d["correlation_id"], []).append(d)

    alerts, seen = [], set()
    for job in jobs:
        cid = job["correlation_id"]
        if cid in seen:
            continue
        seen.add(cid)
        alerts.append((job, _read_result(cid), deliveries_by_cid.get(cid, [])))
    for cid in _result_ids_on_disk():
        if cid in seen:
            continue
        result = _read_result(cid)
        if result is None:
            continue
        seen.add(cid)
        alerts.append((None, result, deliveries_by_cid.get(cid, [])))
    return alerts


# ------------------------------------------------------------------ writing

class _SheetWriter:
    def __init__(self) -> None:
        self.truncated_cells = 0

    def cell_value(self, value: Any) -> Any:
        if value is None or isinstance(value, (bool, int, float, datetime)):
            return value
        text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
        text = ILLEGAL_CHARACTERS_RE.sub("�", text)
        if len(text) > EXCEL_CELL_MAX:
            self.truncated_cells += 1
            note = f"… [truncated: Excel cell limit, {len(text)} characters in full]"
            text = text[:EXCEL_CELL_MAX - len(note)] + note
        return text

    def write(self, ws, columns: list[str], rows: list[dict[str, Any]]) -> None:
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r, row in enumerate(rows, start=2):
            for c, column in enumerate(columns, start=1):
                value = self.cell_value(row.get(column))
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"           # text, never a formula
                    if "\n" in value:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                elif isinstance(value, datetime):
                    cell.number_format = _DATETIME_FORMAT
        ws.freeze_panes = "B2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for c, column in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(c)].width = min(max(len(column) + 2, 14), 60)


def _columns(rows: list[dict[str, Any]], leading: tuple = (), trailing: tuple = ()) -> list[str]:
    present = {key for row in rows for key in row}
    ordered = [c for c in leading if c in present or not rows]
    for row in rows:
        for key in row:
            if key not in ordered and key not in trailing:
                ordered.append(key)
    return ordered + [c for c in trailing if c in present or not rows]


def _build_workbook(alerts: list[tuple]) -> tuple[bytes, int]:
    alert_rows = [_alert_row(job, result, dl) for job, result, dl in alerts]
    cids = [row["correlation_id"] for row in alert_rows]
    factor_rows = [r for (job, result, _), cid in zip(alerts, cids) for r in _risk_factor_rows(cid, result)]
    email_rows = [r for (job, result, dl), cid in zip(alerts, cids) for r in _email_rows(cid, result, dl)]

    writer = _SheetWriter()
    wb = Workbook()
    ws = wb.active
    ws.title = "Alerts"
    writer.write(ws, _columns(alert_rows, _LEADING_COLUMNS, _TRAILING_COLUMNS), alert_rows)
    writer.write(wb.create_sheet("Risk factors"),
                 ["correlation_id", "order", "rule", "effect", "detail"], factor_rows)
    writer.write(wb.create_sheet("Emails"), _columns(email_rows) or [
        "correlation_id", "email_id", "notification_type", "kind", "recipients", "subject",
        "delivery_status", "attempts", "remote_id", "error", "created_at (JST)", "updated_at (JST)",
    ], email_rows)

    records = [row["record"] for row in alert_rows]
    info = wb.create_sheet("Export info")
    for label, value in (
        ("Exported at (JST)", datetime.now(JST).replace(tzinfo=None, microsecond=0)),
        ("Alerts", len(alert_rows)),
        (f"  {RECORD_COMPLETE}", records.count(RECORD_COMPLETE)),
        (f"  {RECORD_IN_PROGRESS}", records.count(RECORD_IN_PROGRESS)),
        (f"  {RECORD_RESULT_ONLY}", records.count(RECORD_RESULT_ONLY)),
        ("Risk factor rows", len(factor_rows)),
        ("Email rows", len(email_rows)),
        ("Cells truncated at Excel's 32,767-character limit", writer.truncated_cells),
        ("Source", "jobs table + output/alerts result files + email_deliveries table; "
                   "complete history, one row per correlation_id"),
        ("Times", "Columns marked (JST) are Japan Standard Time (UTC+09:00)"),
    ):
        info.append([label, value])
        if isinstance(value, datetime):
            info.cell(row=info.max_row, column=2).number_format = _DATETIME_FORMAT
    info.column_dimensions["A"].width = 52
    info.column_dimensions["B"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(alert_rows)


async def export_alerts_xlsx() -> tuple[bytes, int]:
    """(workbook bytes, number of alerts in it)."""
    jobs = await job_store.list_all_jobs()
    deliveries = await delivery_store.list_all_deliveries()
    # File reads and workbook assembly are blocking; keep them off the event loop.
    alerts = await asyncio.to_thread(_collect, jobs, deliveries)
    content, count = await asyncio.to_thread(_build_workbook, alerts)
    logger.info("alerts_exported", alerts=count, bytes=len(content))
    return content, count
